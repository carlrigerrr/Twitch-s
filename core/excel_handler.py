import os
from datetime import datetime, timedelta
from config import EXCEL_CLEANUP_DAYS

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

class ExcelHandler:
    def __init__(self):
        self.workbook = None
        self.worksheet = None
        self.excel_path = ""
        
    def is_available(self):
        """Check if openpyxl is available"""
        return OPENPYXL_AVAILABLE
    
    def create_new_excel(self, filepath):
        """Create a new Excel file"""
        if not OPENPYXL_AVAILABLE:
            return False
            
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "TwitchClips"
            
            # Add headers
            headers = ["Video URL", "Video Title", "Game Category", "Broadcaster", 
                      "Views", "Duration", "Download Date", "Clip ID"]
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                # Style the header
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Adjust column widths
            column_widths = [40, 50, 20, 25, 10, 10, 20, 20]
            for i, width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = width
            
            # Save the file
            wb.save(filepath)
            wb.close()
            
            self.excel_path = filepath
            return True
            
        except Exception:
            return False
    
    def load_excel_file(self, filepath):
        """Load Excel file and clean old data"""
        if not OPENPYXL_AVAILABLE:
            return False, "openpyxl not installed"
        
        try:
            # Load workbook
            self.workbook = openpyxl.load_workbook(filepath)
            self.excel_path = filepath
            
            # Get or create worksheet
            if "TwitchClips" in self.workbook.sheetnames:
                self.worksheet = self.workbook["TwitchClips"]
            else:
                self.worksheet = self.workbook.create_sheet("TwitchClips")
                # Add headers if new sheet
                headers = ["Video URL", "Video Title", "Game Category", "Broadcaster", 
                          "Views", "Duration", "Download Date", "Clip ID"]
                for col, header in enumerate(headers, 1):
                    self.worksheet.cell(row=1, column=col, value=header)
            
            # Clean old data
            removed_count = self.clean_old_data()
            
            # Count remaining records
            record_count = self.worksheet.max_row - 1  # Subtract header row
            
            return True, f"Loaded! {record_count} records (removed {removed_count} old entries)"
            
        except Exception as e:
            return False, f"Failed to load Excel: {str(e)}"
    
    def clean_old_data(self):
        """Remove data older than 30 days from Excel"""
        if not self.worksheet:
            return 0
            
        try:
            cutoff_date = datetime.now() - timedelta(days=EXCEL_CLEANUP_DAYS)
            rows_to_delete = []
            
            # Check each row (skip header)
            for row in range(2, self.worksheet.max_row + 1):
                date_cell = self.worksheet.cell(row=row, column=7)  # Download Date column
                if date_cell.value:
                    try:
                        # Parse the date
                        if isinstance(date_cell.value, datetime):
                            download_date = date_cell.value
                        else:
                            download_date = datetime.strptime(str(date_cell.value), '%Y-%m-%d %H:%M:%S')
                        
                        # Check if older than cutoff
                        if download_date < cutoff_date:
                            rows_to_delete.append(row)
                    except:
                        pass
            
            # Delete old rows (from bottom to top to avoid index issues)
            if rows_to_delete:
                for row in reversed(rows_to_delete):
                    self.worksheet.delete_rows(row)
                
                # Save the workbook
                self.workbook.save(self.excel_path)
                
            return len(rows_to_delete)
            
        except Exception:
            return 0
    
    def check_if_downloaded(self, clip_id):
        """Check if a clip was already downloaded"""
        if not self.worksheet:
            return False
            
        try:
            # Check each row for the clip ID
            for row in range(2, self.worksheet.max_row + 1):
                id_cell = self.worksheet.cell(row=row, column=8)  # Clip ID column
                if id_cell.value == clip_id:
                    return True
            return False
            
        except Exception:
            return False
    
    def add_clip(self, clip_data, game_name):
        """Add downloaded clip info to Excel"""
        if not self.worksheet:
            return False
            
        try:
            # Find next empty row
            next_row = self.worksheet.max_row + 1
            
            # Add data
            data = [
                clip_data['url'],
                clip_data['title'],
                game_name,
                clip_data['broadcaster_name'],
                clip_data['view_count'],
                int(clip_data['duration']),
                datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                clip_data['id']
            ]
            
            for col, value in enumerate(data, 1):
                self.worksheet.cell(row=next_row, column=col, value=value)
            
            # Save the workbook
            self.workbook.save(self.excel_path)
            return True
            
        except Exception:
            return False
    
    def close(self):
        """Close the workbook"""
        if self.workbook:
            try:
                self.workbook.close()
            except:
                pass